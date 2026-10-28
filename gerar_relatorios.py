from pathlib import Path
from datetime import datetime

import pandas as pd
import matplotlib.pyplot as plt
from sqlalchemy import text
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.drawing.image import Image

from database import engine


# ==========================================
# CONFIGURAÇÕES
# ==========================================

PASTA_RELATORIOS = Path('Relatorios')


# ==========================================
# CONSULTAS SQL
# ==========================================

def consultar_resumo():
    """Retorna os indicadores gerais das vendas."""

    consulta = text("""
        SELECT
            COUNT(DISTINCT (id_loja, codigo_venda)) AS quantidade_pedidos,
            COALESCE(SUM(quantidade), 0) AS unidades_vendidas,
            COALESCE(SUM(valor_final), 0) AS faturamento_total,
            COALESCE(
                ROUND(
                    SUM(valor_final) / NULLIF(COUNT(DISTINCT (id_loja, codigo_venda)), 0),
                    2
                ),
                0
            ) AS ticket_medio
        FROM vendas
    """)

    with engine.connect() as conexao:
        return pd.read_sql_query(consulta, conexao)


def consultar_faturamento_lojas():
    """Retorna o faturamento de cada loja."""

    consulta = text("""
        SELECT
            l.nome AS loja,
            COUNT(DISTINCT (v.id_loja, v.codigo_venda)) AS quantidade_pedidos,
            COALESCE(SUM(v.quantidade), 0) AS unidades_vendidas,
            COALESCE(SUM(v.valor_final), 0) AS faturamento
        FROM lojas l
        LEFT JOIN vendas v ON l.id_loja = v.id_loja
        GROUP BY l.id_loja, l.nome
        ORDER BY faturamento DESC
    """)

    with engine.connect() as conexao:
        return pd.read_sql_query(consulta, conexao)


def consultar_faturamento_mensal():
    """Retorna o faturamento agrupado por mês."""

    consulta = text("""
        SELECT
            DATE_TRUNC('month', data_venda)::date AS mes,
            COUNT(DISTINCT (id_loja, codigo_venda)) AS quantidade_pedidos,
            COALESCE(SUM(quantidade), 0) AS unidades_vendidas,
            COALESCE(SUM(valor_final), 0) AS faturamento
        FROM vendas
        GROUP BY DATE_TRUNC('month', data_venda)
        ORDER BY mes
    """)

    with engine.connect() as conexao:
        return pd.read_sql_query(consulta, conexao)


def consultar_produtos():
    """Retorna o desempenho de cada produto."""

    consulta = text("""
        SELECT
            produto,
            SUM(quantidade) AS unidades_vendidas,
            COUNT(DISTINCT (id_loja, codigo_venda)) AS pedidos_com_produto,
            SUM(valor_final) AS faturamento
        FROM vendas
        GROUP BY produto
        ORDER BY faturamento DESC
    """)

    with engine.connect() as conexao:
        return pd.read_sql_query(consulta, conexao)


# ==========================================
# GRÁFICOS
# ==========================================

def criar_graficos(faturamento_lojas, faturamento_mensal, produtos):
    """Gera gráficos PNG temporários para inserir no Excel."""

    PASTA_RELATORIOS.mkdir(parents=True, exist_ok=True)
    caminhos = {}

    if not faturamento_lojas.empty:
        caminho = PASTA_RELATORIOS / 'grafico_faturamento_lojas.png'
        fig, ax = plt.subplots(figsize=(9, 5))

        dados = faturamento_lojas.sort_values(
            'faturamento',
            ascending=True,
        )

        ax.barh(dados['loja'], dados['faturamento'])
        ax.set_title('Faturamento por loja')
        ax.set_xlabel('Faturamento (R$)')
        ax.grid(axis='x', alpha=0.25)

        fig.tight_layout()
        fig.savefig(caminho, dpi=150, bbox_inches='tight')
        plt.close(fig)

        caminhos['lojas'] = caminho

    if not faturamento_mensal.empty:
        caminho = PASTA_RELATORIOS / 'grafico_faturamento_mensal.png'
        fig, ax = plt.subplots(figsize=(9, 5))

        dados = faturamento_mensal.copy()
        dados['mes'] = pd.to_datetime(dados['mes'])

        ax.plot(dados['mes'], dados['faturamento'], marker='o')
        ax.set_title('Evolução mensal do faturamento')
        ax.set_xlabel('Mês')
        ax.set_ylabel('Faturamento (R$)')
        ax.grid(alpha=0.25)

        fig.autofmt_xdate()
        fig.tight_layout()
        fig.savefig(caminho, dpi=150, bbox_inches='tight')
        plt.close(fig)

        caminhos['mensal'] = caminho

    if not produtos.empty:
        caminho = PASTA_RELATORIOS / 'grafico_top_10_produtos.png'

        dados = produtos.head(10).sort_values(
            'faturamento',
            ascending=True,
        )

        fig, ax = plt.subplots(figsize=(10, 6))
        ax.barh(dados['produto'].astype(str), dados['faturamento'])
        ax.set_title('Top 10 produtos por faturamento')
        ax.set_xlabel('Faturamento (R$)')
        ax.grid(axis='x', alpha=0.25)

        fig.tight_layout()
        fig.savefig(caminho, dpi=150, bbox_inches='tight')
        plt.close(fig)

        caminhos['produtos'] = caminho

    return caminhos


# ==========================================
# FORMATAÇÃO DO EXCEL
# ==========================================

def formatar_planilha(planilha, nome_aba):
    """Aplica formatação básica e formatos numéricos."""

    preenchimento = PatternFill('solid', fgColor='1F4E78')
    fonte = Font(color='FFFFFF', bold=True)

    for celula in planilha[1]:
        celula.fill = preenchimento
        celula.font = fonte
        celula.alignment = Alignment(
            horizontal='center',
            vertical='center',
        )

    planilha.freeze_panes = 'A2'

    if planilha.max_row >= 2 and planilha.max_column >= 1:
        planilha.auto_filter.ref = planilha.dimensions

    for coluna in planilha.columns:
        maior_largura = max(
            len(str(celula.value or ''))
            for celula in coluna
        )

        letra = coluna[0].column_letter
        planilha.column_dimensions[letra].width = min(
            max(maior_largura + 2, 12),
            38,
        )

    # Formata valores monetários pelo nome da coluna.
    for cabecalho in planilha[1]:
        nome_coluna = str(cabecalho.value or '').lower()

        if any(
            palavra in nome_coluna
            for palavra in ('faturamento', 'ticket_medio', 'ticket médio', 'valor_medio_venda', 'valor médio')
        ):
            for linha in range(2, planilha.max_row + 1):
                planilha.cell(
                    linha,
                    cabecalho.column,
                ).number_format = '"R$" #,##0.00'

        if nome_coluna == 'mes':
            for linha in range(2, planilha.max_row + 1):
                planilha.cell(
                    linha,
                    cabecalho.column,
                ).number_format = 'MM/YYYY'


# ==========================================
# EXPORTAÇÃO PARA EXCEL
# ==========================================

def exportar_excel(resumo, faturamento_lojas, faturamento_mensal, produtos):
    """Exporta os relatórios e gráficos para um arquivo Excel."""

    PASTA_RELATORIOS.mkdir(parents=True, exist_ok=True)

    data_hora = datetime.now().strftime('%Y%m%d_%H%M%S')
    arquivo_saida = PASTA_RELATORIOS / f'relatorio_vendas_{data_hora}.xlsx'

    graficos = criar_graficos(
        faturamento_lojas,
        faturamento_mensal,
        produtos,
    )

    with pd.ExcelWriter(
        arquivo_saida,
        engine='openpyxl',
        datetime_format='DD/MM/YYYY',
    ) as escritor:

        resumo.to_excel(
            escritor,
            sheet_name='Resumo Geral',
            index=False,
        )

        faturamento_lojas.to_excel(
            escritor,
            sheet_name='Faturamento por Loja',
            index=False,
        )

        faturamento_mensal.to_excel(
            escritor,
            sheet_name='Faturamento Mensal',
            index=False,
        )

        produtos.to_excel(
            escritor,
            sheet_name='Produtos',
            index=False,
        )

        for nome_aba, planilha in escritor.sheets.items():
            formatar_planilha(planilha, nome_aba)

        # Adiciona título e data de geração ao resumo.
        planilha_resumo = escritor.sheets['Resumo Geral']
        planilha_resumo.insert_rows(1, 3)
        planilha_resumo.freeze_panes = 'A5'
        planilha_resumo.auto_filter.ref = (
            f'A4:{planilha_resumo.cell(4, planilha_resumo.max_column).column_letter}'
            f'{planilha_resumo.max_row}'
        )

        planilha_resumo['A1'] = 'RELATÓRIO DE VENDAS'
        planilha_resumo['A1'].font = Font(
            size=18,
            bold=True,
            color='FFFFFF',
        )
        planilha_resumo['A1'].fill = PatternFill(
            'solid',
            fgColor='17365D',
        )
        planilha_resumo.merge_cells('A1:D1')

        planilha_resumo['A2'] = (
            f'Gerado em: {datetime.now():%d/%m/%Y às %H:%M}'
        )
        planilha_resumo.merge_cells('A2:D2')

        for celula in planilha_resumo[4]:
            celula.fill = PatternFill('solid', fgColor='1F4E78')
            celula.font = Font(color='FFFFFF', bold=True)

        # Gráfico de faturamento por loja.
        if 'lojas' in graficos:
            aba = escritor.sheets['Faturamento por Loja']

            grafico = BarChart()
            grafico.type = 'bar'
            grafico.style = 10
            grafico.title = 'Faturamento por loja'
            grafico.y_axis.title = 'Loja'
            grafico.x_axis.title = 'Faturamento (R$)'
            grafico.height = 8
            grafico.width = 15

            dados = Reference(
                aba,
                min_col=4,
                min_row=1,
                max_row=aba.max_row,
            )
            categorias = Reference(
                aba,
                min_col=1,
                min_row=2,
                max_row=aba.max_row,
            )

            grafico.add_data(dados, titles_from_data=True)
            grafico.set_categories(categorias)
            aba.add_chart(grafico, 'F2')

        # Gráfico de evolução mensal.
        if 'mensal' in graficos:
            aba = escritor.sheets['Faturamento Mensal']

            grafico = LineChart()
            grafico.style = 13
            grafico.title = 'Evolução mensal do faturamento'
            grafico.y_axis.title = 'Faturamento (R$)'
            grafico.x_axis.title = 'Mês'
            grafico.height = 8
            grafico.width = 15

            dados = Reference(
                aba,
                min_col=4,
                min_row=1,
                max_row=aba.max_row,
            )
            categorias = Reference(
                aba,
                min_col=1,
                min_row=2,
                max_row=aba.max_row,
            )

            grafico.add_data(dados, titles_from_data=True)
            grafico.set_categories(categorias)
            aba.add_chart(grafico, 'F2')

        # Gráfico dos dez produtos com maior faturamento.
        if 'produtos' in graficos:
            aba = escritor.sheets['Produtos']

            grafico = BarChart()
            grafico.type = 'bar'
            grafico.style = 10
            grafico.title = 'Top 10 produtos por faturamento'
            grafico.y_axis.title = 'Produto'
            grafico.x_axis.title = 'Faturamento (R$)'
            grafico.height = 10
            grafico.width = 16

            max_linha = min(aba.max_row, 11)

            dados = Reference(
                aba,
                min_col=4,
                min_row=1,
                max_row=max_linha,
            )
            categorias = Reference(
                aba,
                min_col=1,
                min_row=2,
                max_row=max_linha,
            )

            grafico.add_data(dados, titles_from_data=True)
            grafico.set_categories(categorias)
            aba.add_chart(grafico, 'F2')

        # Insere uma imagem do gráfico de lojas no resumo.
        if 'lojas' in graficos:
            imagem = Image(str(graficos['lojas']))
            imagem.width = 600
            imagem.height = 330
            planilha_resumo.add_image(imagem, 'A7')

    # Remove os arquivos PNG temporários após salvar o Excel.
    for caminho in graficos.values():
        try:
            caminho.unlink(missing_ok=True)
        except OSError:
            pass

    print('\nRelatório gerado com sucesso!')
    print(f'Arquivo: {arquivo_saida.resolve()}')

    return arquivo_saida


# ==========================================
# EXECUÇÃO PRINCIPAL
# ==========================================

def main():
    print('=' * 45)
    print('GERADOR DE RELATÓRIOS DE VENDAS')
    print('=' * 45)

    print('\nConsultando o banco de dados...')

    resumo = consultar_resumo()
    faturamento_lojas = consultar_faturamento_lojas()
    faturamento_mensal = consultar_faturamento_mensal()
    produtos = consultar_produtos()

    print('Consultas concluídas!')

    faturamento = float(resumo.loc[0, 'faturamento_total'])

    print(f'\nFaturamento total: R$ {faturamento:,.2f}')
    print(f'Lojas encontradas: {len(faturamento_lojas)}')
    quantidade_pedidos = int(resumo.loc[0, 'quantidade_pedidos'])
    unidades_vendidas = int(resumo.loc[0, 'unidades_vendidas'])
    ticket_medio = float(resumo.loc[0, 'ticket_medio'])

    print(f'Pedidos encontrados: {quantidade_pedidos:,}'.replace(',', '.'))
    print(f'Unidades vendidas: {unidades_vendidas:,}'.replace(',', '.'))
    print(f'Ticket médio por pedido: R$ {ticket_medio:,.2f}')
    print(f'Produtos encontrados: {len(produtos)}')

    exportar_excel(
        resumo,
        faturamento_lojas,
        faturamento_mensal,
        produtos,
    )


if __name__ == '__main__':
    main()
