from enviar_relatorios import fmt_moeda, fmt_num
from email import message_from_bytes
from email.policy import default
from unittest.mock import patch
from unittest.mock import MagicMock
import enviar_relatorios
from openpyxl import load_workbook

def test_fmt_moeda():
    assert fmt_moeda(1234.56) == 'R$ 1.234,56'


def test_fmt_moeda_valor_zero():
    assert fmt_moeda(0) == 'R$ 0,00'


def test_fmt_num():
    assert fmt_num(1234) == '1.234'


def test_fmt_num_valor_zero():
    assert fmt_num(0) == '0'
from unittest.mock import patch

import enviar_relatorios


def test_enviar_email_modo_simulacao():
    with patch.object(enviar_relatorios, 'ENVIAR_EMAILS', False):
        resultado = enviar_relatorios.enviar_email(
            destino='gerente@exemplo.com',
            assunto='Relatório diário',
            corpo='Este é um teste.',
        )

    assert resultado is False


def test_enviar_email_modo_simulacao_nao_conecta_smtp():
    with patch.object(enviar_relatorios, 'ENVIAR_EMAILS', False):
        with patch.object(enviar_relatorios.smtplib, 'SMTP') as smtp_mock:
            resultado = enviar_relatorios.enviar_email(
                destino='gerente@exemplo.com',
                assunto='Relatório diário',
                corpo='Este é um teste.',
            )

    assert resultado is False
    smtp_mock.assert_not_called()
from unittest.mock import patch

from enviar_relatorios import enviar_email


def test_enviar_email_simulacao_com_anexo(tmp_path):
    arquivo = tmp_path / 'relatorio.xlsx'
    arquivo.write_bytes(b'conteudo de teste')

    with patch('enviar_relatorios.ENVIAR_EMAILS', False):
        resultado = enviar_email(
            destino='gerente@exemplo.com',
            assunto='Relatório diário',
            corpo='Segue o relatório.',
            anexos=[arquivo],
        )

    assert resultado is False
    assert arquivo.exists()



def test_enviar_email_inclui_anexo(tmp_path):
    arquivo = tmp_path / 'relatorio.xlsx'
    arquivo.write_bytes(b'conteudo de teste')

    with (
        patch.object(enviar_relatorios, 'ENVIAR_EMAILS', True),
        patch.object(enviar_relatorios, 'EMAIL_REMETENTE', 'teste@exemplo.com'),
        patch.object(enviar_relatorios, 'SENHA_APP', 'senha_teste'),
        patch.object(enviar_relatorios, 'EMAIL_TESTE', 'destino@exemplo.com'),
        patch.object(enviar_relatorios.smtplib, 'SMTP') as smtp_mock,
    ):
        resultado = enviar_relatorios.enviar_email(
            destino='gerente@exemplo.com',
            assunto='Relatório diário',
            corpo='Segue o relatório.',
            anexos=[arquivo],
        )

    assert resultado is True

    smtp_mock.assert_called_once_with(
        'smtp.gmail.com',
        587,
        timeout=30,
    )

    servidor = smtp_mock.return_value.__enter__.return_value
    servidor.send_message.assert_called_once()

    mensagem = servidor.send_message.call_args.args[0]

    assert mensagem['To'] == 'destino@exemplo.com'
    assert mensagem['Subject'] == '[TESTE] Relatório diário'

    anexos = list(mensagem.iter_attachments())

    assert len(anexos) == 1
    assert anexos[0].get_filename() == 'relatorio.xlsx'
    assert anexos[0].get_payload(decode=True) == b'conteudo de teste'

import smtplib
from unittest.mock import patch

import pytest

import enviar_relatorios


def test_enviar_email_propaga_falha_de_autenticacao():
    with (
        patch.object(enviar_relatorios, 'ENVIAR_EMAILS', True),
        patch.object(enviar_relatorios, 'EMAIL_REMETENTE', 'teste@exemplo.com'),
        patch.object(enviar_relatorios, 'SENHA_APP', 'senha_teste'),
        patch.object(enviar_relatorios, 'EMAIL_TESTE', 'destino@exemplo.com'),
        patch.object(
            enviar_relatorios.smtplib,
            'SMTP',
        ) as smtp_mock,
    ):
        servidor = smtp_mock.return_value.__enter__.return_value
        servidor.login.side_effect = smtplib.SMTPAuthenticationError(
            535,
            b'Autenticacao recusada',
        )

        with pytest.raises(smtplib.SMTPAuthenticationError):
            enviar_relatorios.enviar_email(
                destino='gerente@exemplo.com',
                assunto='Relatório diário',
                corpo='Este é um teste.',
            )

    servidor.login.assert_called_once()
    servidor.send_message.assert_not_called()
def test_enviar_email_tenta_novamente_apos_falha_de_conexao():
    with (
        patch.object(enviar_relatorios, 'ENVIAR_EMAILS', True),
        patch.object(
            enviar_relatorios,
            'EMAIL_REMETENTE',
            'teste@exemplo.com',
        ),
        patch.object(
            enviar_relatorios,
            'SENHA_APP',
            'senha_teste',
        ),
        patch.object(
            enviar_relatorios,
            'EMAIL_TESTE',
            'destino@exemplo.com',
        ),
        patch.object(enviar_relatorios.smtplib, 'SMTP') as smtp_mock,
        patch.object(enviar_relatorios.time, 'sleep') as sleep_mock,
    ):
        servidor = smtp_mock.return_value.__enter__.return_value

        servidor.send_message.side_effect = [
            smtplib.SMTPServerDisconnected('Conexão interrompida'),
            None,
        ]

        resultado = enviar_relatorios.enviar_email(
            destino='gerente@exemplo.com',
            assunto='Relatório diário',
            corpo='Este é um teste.',
        )

    assert resultado is True
    assert smtp_mock.call_count == 2
    assert servidor.send_message.call_count == 2
    sleep_mock.assert_called_once_with(5)
def test_enviar_email_encerra_apos_esgotar_tentativas():
    with (
        patch.object(enviar_relatorios, 'ENVIAR_EMAILS', True),
        patch.object(enviar_relatorios, 'EMAIL_REMETENTE', 'teste@exemplo.com'),
        patch.object(enviar_relatorios, 'SENHA_APP', 'senha_teste'),
        patch.object(enviar_relatorios, 'EMAIL_TESTE', 'destino@exemplo.com'),
        patch.object(enviar_relatorios.smtplib, 'SMTP') as smtp_mock,
        patch.object(enviar_relatorios.time, 'sleep') as sleep_mock,
    ):
        servidor = smtp_mock.return_value.__enter__.return_value
        servidor.send_message.side_effect = smtplib.SMTPServerDisconnected(
            'Conexão interrompida'
        )

        with pytest.raises(smtplib.SMTPServerDisconnected):
            enviar_relatorios.enviar_email(
                destino='gerente@exemplo.com',
                assunto='Relatório diário',
                corpo='Este é um teste.',
            )

    assert servidor.send_message.call_count == 3
    assert sleep_mock.call_count == 2
def test_verificar_login_com_sucesso(monkeypatch):
    servidor_simulado = MagicMock()
    contexto_smtp = MagicMock()
    contexto_smtp.__enter__.return_value = servidor_simulado

    monkeypatch.setattr(
        enviar_relatorios.smtplib,
        'SMTP',
        MagicMock(return_value=contexto_smtp),
    )

    enviar_relatorios.verificar_login()

    servidor_simulado.ehlo.assert_called()
    servidor_simulado.starttls.assert_called_once()
    servidor_simulado.login.assert_called_once_with(
        enviar_relatorios.EMAIL_REMETENTE,
        enviar_relatorios.SENHA_APP,
    )
def test_verificar_login_falha_autenticacao(monkeypatch):
    servidor_simulado = MagicMock()

    servidor_simulado.__enter__.return_value = servidor_simulado
    servidor_simulado.login.side_effect = smtplib.SMTPAuthenticationError(
        535,
        b'Authentication failed',
    )

    monkeypatch.setattr(
        enviar_relatorios.smtplib,
        'SMTP',
        MagicMock(
            return_value=servidor_simulado,
        ),
    )

    with pytest.raises(SystemExit) as erro:
        enviar_relatorios.verificar_login()

    assert erro.value.code == 1


def test_verificar_login_falha_conexao(monkeypatch):
    servidor_simulado = MagicMock()

    servidor_simulado.__enter__.return_value = servidor_simulado
    servidor_simulado.starttls.side_effect = OSError(
        'Falha simulada de conexão',
    )

    monkeypatch.setattr(
        enviar_relatorios.smtplib,
        'SMTP',
        MagicMock(
            return_value=servidor_simulado,
        ),
    )

    with pytest.raises(SystemExit) as erro:
        enviar_relatorios.verificar_login()

    assert erro.value.code == 1
def test_fluxo_completo_postgresql_para_excel(
    engine_teste,
    limpar_tabelas,
    tmp_path,
    monkeypatch,
):
    import gerar_relatorios
    from sqlalchemy import text

    # Direciona as consultas para o banco de testes.
    monkeypatch.setattr(
        gerar_relatorios,
        'engine',
        engine_teste,
    )

    # Direciona os arquivos gerados para uma pasta temporária.
    monkeypatch.setattr(
        gerar_relatorios,
        'PASTA_RELATORIOS',
        tmp_path,
    )

    # Insere uma loja e vendas fictícias no PostgreSQL.
    with engine_teste.begin() as conexao:
        conexao.execute(
            text("""
                INSERT INTO lojas (id_loja, nome, gerente, email)
                VALUES ('1', 'Loja A', 'Maria', 'maria@exemplo.com')
            """)
        )

        conexao.execute(
            text("""
                INSERT INTO vendas (
                    id_loja,
                    data_venda,
                    codigo_venda,
                    produto,
                    quantidade,
                    valor_final
                )
                VALUES
                    (
                        '1',
                        '2026-01-15',
                        '100',
                        'Produto A',
                        2,
                        100.00
                    ),
                    (
                        '1',
                        '2026-01-15',
                        '101',
                        'Produto B',
                        3,
                        200.00
                    )
            """)
        )

    # Executa as consultas reais no banco de testes.
    resumo = gerar_relatorios.consultar_resumo()
    lojas = gerar_relatorios.consultar_faturamento_lojas()
    mensal = gerar_relatorios.consultar_faturamento_mensal()
    produtos = gerar_relatorios.consultar_produtos()

    # Confere os indicadores calculados.
    assert resumo.loc[0, 'quantidade_pedidos'] == 2
    assert resumo.loc[0, 'unidades_vendidas'] == 5
    assert float(resumo.loc[0, 'faturamento_total']) == 300.00
    assert float(resumo.loc[0, 'ticket_medio']) == 150.00

    assert len(lojas) == 1
    assert lojas.loc[0, 'loja'] == 'Loja A'
    assert float(lojas.loc[0, 'faturamento']) == 300.00

    assert len(mensal) == 1
    assert float(mensal.loc[0, 'faturamento']) == 300.00

    assert len(produtos) == 2

    # Gera a planilha Excel com os resultados reais.
    arquivo = gerar_relatorios.exportar_excel(
        resumo,
        lojas,
        mensal,
        produtos,
    )

    # Confere se o arquivo foi criado corretamente.
    assert arquivo.exists()
    assert arquivo.suffix == '.xlsx'

    planilha = load_workbook(arquivo, data_only=True)

    try:
        assert planilha.sheetnames == [
            'Resumo Geral',
            'Faturamento por Loja',
            'Faturamento Mensal',
            'Produtos',
        ]

        assert planilha['Resumo Geral']['A1'].value == (
            'RELATÓRIO DE VENDAS'
        )

        # Confere se os gráficos PNG temporários foram removidos.
        assert not list(tmp_path.glob('*.png'))

    finally:
        planilha.close()
def test_enviar_email_modo_teste(monkeypatch, caplog):
    import logging
    import enviar_relatorios

    monkeypatch.setattr(enviar_relatorios, 'ENVIAR_EMAILS', False)

    with caplog.at_level(logging.INFO):
        resultado = enviar_relatorios.enviar_email(
            destino='teste@exemplo.com',
            assunto='Relatório de vendas',
            corpo='Conteúdo do relatório',
        )

    assert resultado is False
    assert 'E-mail não enviado' in caplog.text
    assert 'teste@exemplo.com' in caplog.text
def test_enviar_email_com_sucesso(monkeypatch):
    import enviar_relatorios

    monkeypatch.setattr(enviar_relatorios, 'ENVIAR_EMAILS', True)
    monkeypatch.setattr(enviar_relatorios, 'EMAIL_REMETENTE', 'remetente@exemplo.com')
    monkeypatch.setattr(enviar_relatorios, 'EMAIL_TESTE', 'teste@exemplo.com')
    monkeypatch.setattr(enviar_relatorios, 'SENHA_APP', 'senha_ficticia')

    class SMTPSimulado:
        def __init__(self, servidor, porta, timeout):
            assert servidor == 'smtp.gmail.com'
            assert porta == 587
            assert timeout == 30

        def __enter__(self):
            return self

        def __exit__(self, tipo, valor, traceback):
            return False

        def ehlo(self):
            pass

        def starttls(self):
            pass

        def login(self, remetente, senha):
            assert remetente == 'remetente@exemplo.com'
            assert senha == 'senha_ficticia'

        def send_message(self, mensagem):
            assert mensagem['To'] == 'teste@exemplo.com'
            assert mensagem['Subject'] == '[TESTE] Relatório de vendas'

    monkeypatch.setattr(
        enviar_relatorios.smtplib,
        'SMTP',
        SMTPSimulado,
    )

    resultado = enviar_relatorios.enviar_email(
        destino='destino@exemplo.com',
        assunto='Relatório de vendas',
        corpo='Conteúdo do relatório',
    )

    assert resultado is True
def test_enviar_email_falha_autenticacao(monkeypatch, caplog):
    import logging
    import smtplib
    import pytest
    import enviar_relatorios

    monkeypatch.setattr(enviar_relatorios, 'ENVIAR_EMAILS', True)
    monkeypatch.setattr(
        enviar_relatorios,
        'EMAIL_REMETENTE',
        'remetente@exemplo.com',
    )
    monkeypatch.setattr(
        enviar_relatorios,
        'EMAIL_TESTE',
        'teste@exemplo.com',
    )
    monkeypatch.setattr(
        enviar_relatorios,
        'SENHA_APP',
        'senha_ficticia',
    )

    class SMTPSimulado:
        def __init__(self, *args, **kwargs):
            pass

        def __enter__(self):
            return self

        def __exit__(self, tipo, valor, traceback):
            return False

        def ehlo(self):
            pass

        def starttls(self):
            pass

        def login(self, remetente, senha):
            raise smtplib.SMTPAuthenticationError(
                535,
                b'Falha de autenticacao simulada',
            )

    monkeypatch.setattr(
        enviar_relatorios.smtplib,
        'SMTP',
        SMTPSimulado,
    )

    with caplog.at_level(logging.ERROR):
        with pytest.raises(smtplib.SMTPAuthenticationError):
            enviar_relatorios.enviar_email(
                destino='destino@exemplo.com',
                assunto='Relatório de vendas',
                corpo='Conteúdo do relatório',
            )

    assert 'Falha de autenticação no Gmail' in caplog.text
def test_importacao_falha_sem_credenciais_de_email(monkeypatch):
    import os
    import subprocess
    import sys

    ambiente = os.environ.copy()
    ambiente['ENVIAR_EMAILS'] = 'true'
    ambiente['EMAIL_REMETENTE'] = ''
    ambiente['SENHA_APP'] = 'senha_teste'
    ambiente['EMAIL_TESTE'] = 'teste@exemplo.com'

    resultado = subprocess.run(
        [
            sys.executable,
            '-c',
            'import enviar_relatorios',
        ],
        env=ambiente,
        capture_output=True,
        text=True,
    )

    assert resultado.returncode != 0
    assert 'EMAIL_REMETENTE e SENHA_APP' in resultado.stderr
def test_importacao_falha_com_email_teste_invalido():
    import os
    import subprocess
    import sys

    ambiente = os.environ.copy()
    ambiente['ENVIAR_EMAILS'] = 'true'
    ambiente['EMAIL_REMETENTE'] = 'remetente@exemplo.com'
    ambiente['SENHA_APP'] = 'senha_teste'
    ambiente['EMAIL_TESTE'] = 'email_invalido'

    resultado = subprocess.run(
        [
            sys.executable,
            '-c',
            'import enviar_relatorios',
        ],
        env=ambiente,
        capture_output=True,
        text=True,
    )

    assert resultado.returncode != 0
    assert 'EMAIL_TESTE' in resultado.stderr
def test_cor_cenario_quando_atinge_meta():
    from enviar_relatorios import cor_cenario

    assert cor_cenario(1000, 1000) == "#00b050"
    assert cor_cenario(1200, 1000) == "#00b050"


def test_cor_cenario_quando_nao_atinge_meta():
    from enviar_relatorios import cor_cenario

    assert cor_cenario(999, 1000) == "#e53935"


def test_fmt_moeda():
    from enviar_relatorios import fmt_moeda

    assert fmt_moeda(1234.56) == "R$ 1.234,56"
    assert fmt_moeda(1000) == "R$ 1.000,00"
    assert fmt_moeda(0) == "R$ 0,00"


def test_fmt_num():
    from enviar_relatorios import fmt_num

    assert fmt_num(1234) == "1.234"
    assert fmt_num(1000000) == "1.000.000"
    assert fmt_num(0) == "0"
def test_ler_arquivo_csv_valido(tmp_path):
    import pandas as pd
    from enviar_relatorios import ler_arquivo

    arquivo = tmp_path / "dados.csv"
    pd.DataFrame({
        "Loja": ["Loja A"],
        "Valor": [150.50],
    }).to_csv(arquivo, index=False)

    resultado = ler_arquivo(
        arquivo,
        ["Loja", "Valor"],
    )

    assert len(resultado) == 1
    assert resultado.loc[0, "Loja"] == "Loja A"
    assert resultado.loc[0, "Valor"] == 150.50


def test_ler_arquivo_arquivo_inexistente(tmp_path):
    import pytest
    from enviar_relatorios import ler_arquivo

    arquivo = tmp_path / "inexistente.xlsx"

    with pytest.raises(FileNotFoundError, match="Arquivo não encontrado"):
        ler_arquivo(arquivo, ["Loja"])


def test_ler_arquivo_colunas_obrigatorias_ausentes(tmp_path):
    import pandas as pd
    import pytest
    from enviar_relatorios import ler_arquivo

    arquivo = tmp_path / "dados.xlsx"
    pd.DataFrame({
        "Loja": ["Loja A"],
    }).to_excel(arquivo, index=False)

    with pytest.raises(ValueError, match="Colunas ausentes"):
        ler_arquivo(
            arquivo,
            ["Loja", "Gerente", "E-mail"],
        )
def test_ler_arquivo_csv_valido(tmp_path):
    import pandas as pd
    from enviar_relatorios import ler_arquivo

    arquivo = tmp_path / "dados.csv"
    pd.DataFrame({
        "Loja": ["Loja A"],
        "Valor": [150.50],
    }).to_csv(arquivo, index=False)

    resultado = ler_arquivo(
        arquivo,
        ["Loja", "Valor"],
    )

    assert len(resultado) == 1
    assert resultado.loc[0, "Loja"] == "Loja A"
    assert resultado.loc[0, "Valor"] == 150.50


def test_ler_arquivo_arquivo_inexistente(tmp_path):
    import pytest
    from enviar_relatorios import ler_arquivo

    arquivo = tmp_path / "inexistente.xlsx"

    with pytest.raises(FileNotFoundError, match="Arquivo não encontrado"):
        ler_arquivo(arquivo, ["Loja"])


def test_ler_arquivo_colunas_obrigatorias_ausentes(tmp_path):
    import pandas as pd
    import pytest
    from enviar_relatorios import ler_arquivo

    arquivo = tmp_path / "dados.xlsx"
    pd.DataFrame({
        "Loja": ["Loja A"],
    }).to_excel(arquivo, index=False)

    with pytest.raises(ValueError, match="Colunas ausentes"):
        ler_arquivo(
            arquivo,
            ["Loja", "Gerente", "E-mail"],
        )