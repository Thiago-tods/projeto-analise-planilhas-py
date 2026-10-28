import pandas as pd
import matplotlib
from gerar_relatorios import main
matplotlib.use('Agg')

from pathlib import Path

from openpyxl import load_workbook

from gerar_relatorios import criar_graficos, exportar_excel
from gerar_relatorios import consultar_resumo


def test_consultar_resumo_retorna_indicadores(monkeypatch):
    dados_esperados = pd.DataFrame([{
        'quantidade_pedidos': 3,
        'unidades_vendidas': 4,
        'faturamento_total': 1300.00,
        'ticket_medio': 433.33,
    }])

    def consulta_simulada(consulta, conexao):
        return dados_esperados

    monkeypatch.setattr(
        'gerar_relatorios.pd.read_sql_query',
        consulta_simulada,
    )

    resultado = consultar_resumo()

    assert resultado.loc[0, 'quantidade_pedidos'] == 3
    assert resultado.loc[0, 'unidades_vendidas'] == 4
    assert resultado.loc[0, 'faturamento_total'] == 1300.00
    assert resultado.loc[0, 'ticket_medio'] == 433.33
import pandas as pd

from gerar_relatorios import (
    consultar_faturamento_lojas,
    consultar_faturamento_mensal,
    consultar_produtos,
)


def test_consultar_faturamento_lojas(monkeypatch):
    dados_esperados = pd.DataFrame({
        'loja': ['Loja A', 'Loja B'],
        'faturamento_total': [1000.00, 2000.00],
    })

    def consulta_simulada(consulta, conexao):
        return dados_esperados

    monkeypatch.setattr(
        'gerar_relatorios.pd.read_sql_query',
        consulta_simulada,
    )

    resultado = consultar_faturamento_lojas()

    assert len(resultado) == 2
    assert resultado.loc[0, 'loja'] == 'Loja A'
    assert resultado.loc[1, 'faturamento_total'] == 2000.00


def test_consultar_faturamento_mensal(monkeypatch):
    dados_esperados = pd.DataFrame({
        'mes': ['2026-01', '2026-02'],
        'faturamento_total': [1500.00, 2500.00],
    })

    def consulta_simulada(consulta, conexao):
        return dados_esperados

    monkeypatch.setattr(
        'gerar_relatorios.pd.read_sql_query',
        consulta_simulada,
    )

    resultado = consultar_faturamento_mensal()

    assert len(resultado) == 2
    assert resultado.loc[0, 'mes'] == '2026-01'
    assert resultado.loc[1, 'faturamento_total'] == 2500.00


def test_consultar_produtos(monkeypatch):
    dados_esperados = pd.DataFrame({
        'produto': ['Produto A', 'Produto B'],
        'quantidade_vendida': [10, 20],
        'faturamento_total': [500.00, 1000.00],
    })

    def consulta_simulada(consulta, conexao):
        return dados_esperados

    monkeypatch.setattr(
        'gerar_relatorios.pd.read_sql_query',
        consulta_simulada,
    )

    resultado = consultar_produtos()

    assert len(resultado) == 2
    assert resultado.loc[0, 'produto'] == 'Produto A'
    assert resultado.loc[1, 'quantidade_vendida'] == 20
def test_consultar_faturamento_lojas(monkeypatch):
    dados_esperados = pd.DataFrame({
        'loja': ['Loja A', 'Loja B'],
        'quantidade_pedidos': [2, 3],
        'unidades_vendidas': [10, 20],
        'faturamento': [1000.00, 2000.00],
    })

    def consulta_simulada(consulta, conexao):
        return dados_esperados

    monkeypatch.setattr(
        'gerar_relatorios.pd.read_sql_query',
        consulta_simulada,
    )

    resultado = consultar_faturamento_lojas()

    assert len(resultado) == 2
    assert resultado.loc[0, 'loja'] == 'Loja A'
    assert resultado.loc[0, 'faturamento'] == 1000.00
    assert resultado.loc[1, 'loja'] == 'Loja B'
    assert resultado.loc[1, 'faturamento'] == 2000.00
def test_consultar_faturamento_mensal(monkeypatch):
    dados_esperados = pd.DataFrame({
        'mes': ['2026-01-01', '2026-02-01'],
        'quantidade_pedidos': [3, 5],
        'unidades_vendidas': [15, 25],
        'faturamento': [1500.00, 2500.00],
    })

    def consulta_simulada(consulta, conexao):
        return dados_esperados

    monkeypatch.setattr(
        'gerar_relatorios.pd.read_sql_query',
        consulta_simulada,
    )

    resultado = consultar_faturamento_mensal()

    assert len(resultado) == 2
    assert resultado.loc[0, 'mes'] == '2026-01-01'
    assert resultado.loc[1, 'mes'] == '2026-02-01'
    assert resultado.loc[0, 'faturamento'] == 1500.00
    assert resultado.loc[1, 'faturamento'] == 2500.00
def test_consultar_produtos(monkeypatch):
    dados_esperados = pd.DataFrame({
        'produto': ['Produto A', 'Produto B'],
        'unidades_vendidas': [10, 20],
        'pedidos_com_produto': [4, 7],
        'faturamento': [500.00, 1000.00],
    })

    def consulta_simulada(consulta, conexao):
        return dados_esperados

    monkeypatch.setattr(
        'gerar_relatorios.pd.read_sql_query',
        consulta_simulada,
    )

    resultado = consultar_produtos()

    assert len(resultado) == 2
    assert resultado.loc[0, 'produto'] == 'Produto A'
    assert resultado.loc[1, 'produto'] == 'Produto B'
    assert resultado.loc[0, 'unidades_vendidas'] == 10
    assert resultado.loc[1, 'pedidos_com_produto'] == 7
    assert resultado.loc[1, 'faturamento'] == 1000.00
def test_criar_graficos_gera_tres_arquivos(tmp_path, monkeypatch):
    import gerar_relatorios

    monkeypatch.setattr(
        gerar_relatorios,
        'PASTA_RELATORIOS',
        tmp_path,
    )

    faturamento_lojas = pd.DataFrame({
        'loja': ['Loja A', 'Loja B'],
        'faturamento': [1000.00, 2000.00],
    })

    faturamento_mensal = pd.DataFrame({
        'mes': ['2026-01-01', '2026-02-01'],
        'faturamento': [1500.00, 2500.00],
    })

    produtos = pd.DataFrame({
        'produto': ['Produto A', 'Produto B'],
        'faturamento': [500.00, 1000.00],
    })

    resultado = criar_graficos(
        faturamento_lojas,
        faturamento_mensal,
        produtos,
    )

    assert set(resultado) == {'lojas', 'mensal', 'produtos'}

    for caminho in resultado.values():
        assert caminho.exists()
        assert caminho.stat().st_size > 0
def test_criar_graficos_com_dados_vazios(tmp_path, monkeypatch):
    import gerar_relatorios

    monkeypatch.setattr(
        gerar_relatorios,
        'PASTA_RELATORIOS',
        tmp_path,
    )

    resultado = criar_graficos(
        pd.DataFrame(),
        pd.DataFrame(),
        pd.DataFrame(),
    )

    assert resultado == {}
    assert list(tmp_path.iterdir()) == []
def test_exportar_excel_cria_planilha_com_abas(
    tmp_path,
    monkeypatch,
):
    import gerar_relatorios

    monkeypatch.setattr(
        gerar_relatorios,
        'PASTA_RELATORIOS',
        tmp_path,
    )

    resumo = pd.DataFrame([{
        'quantidade_pedidos': 3,
        'unidades_vendidas': 4,
        'faturamento_total': 1300.00,
        'ticket_medio': 433.33,
    }])

    faturamento_lojas = pd.DataFrame({
        'loja': ['Loja A', 'Loja B'],
        'quantidade_pedidos': [2, 1],
        'unidades_vendidas': [3, 1],
        'faturamento': [1000.00, 300.00],
    })

    faturamento_mensal = pd.DataFrame({
        'mes': pd.to_datetime(['2026-01-01', '2026-02-01']),
        'quantidade_pedidos': [2, 1],
        'unidades_vendidas': [3, 1],
        'faturamento': [1000.00, 300.00],
    })

    produtos = pd.DataFrame({
        'produto': ['Produto A', 'Produto B'],
        'unidades_vendidas': [3, 1],
        'pedidos_com_produto': [2, 1],
        'faturamento': [1000.00, 300.00],
    })

    caminho_arquivo = exportar_excel(
        resumo,
        faturamento_lojas,
        faturamento_mensal,
        produtos,
    )

    assert caminho_arquivo.exists()
    assert caminho_arquivo.suffix == '.xlsx'

    planilha = load_workbook(caminho_arquivo)

    assert planilha.sheetnames == [
        'Resumo Geral',
        'Faturamento por Loja',
        'Faturamento Mensal',
        'Produtos',
    ]

    assert planilha['Resumo Geral']['A1'].value == 'RELATÓRIO DE VENDAS'

    planilha.close()

    # Os gráficos PNG temporários devem ter sido removidos.
    assert not list(tmp_path.glob('*.png'))
def test_main_executa_consultas_e_exporta_relatorio(monkeypatch):
    import gerar_relatorios

    chamadas = []

    resumo = pd.DataFrame([{
        'quantidade_pedidos': 3,
        'unidades_vendidas': 4,
        'faturamento_total': 1300.00,
        'ticket_medio': 433.33,
    }])

    faturamento_lojas = pd.DataFrame({
        'loja': ['Loja A'],
        'faturamento': [1300.00],
    })

    faturamento_mensal = pd.DataFrame({
        'mes': ['2026-01-01'],
        'faturamento': [1300.00],
    })

    produtos = pd.DataFrame({
        'produto': ['Produto A'],
        'unidades_vendidas': [4],
        'pedidos_com_produto': [3],
        'faturamento': [1300.00],
    })

    def consultar_resumo_simulado():
        chamadas.append('resumo')
        return resumo

    def consultar_lojas_simulado():
        chamadas.append('lojas')
        return faturamento_lojas

    def consultar_mensal_simulado():
        chamadas.append('mensal')
        return faturamento_mensal

    def consultar_produtos_simulado():
        chamadas.append('produtos')
        return produtos

    def exportar_excel_simulado(
        resumo_recebido,
        lojas_recebidas,
        mensal_recebido,
        produtos_recebidos,
    ):
        chamadas.append('exportar')

        assert resumo_recebido.equals(resumo)
        assert lojas_recebidas.equals(faturamento_lojas)
        assert mensal_recebido.equals(faturamento_mensal)
        assert produtos_recebidos.equals(produtos)

    monkeypatch.setattr(
        gerar_relatorios,
        'consultar_resumo',
        consultar_resumo_simulado,
    )
    monkeypatch.setattr(
        gerar_relatorios,
        'consultar_faturamento_lojas',
        consultar_lojas_simulado,
    )
    monkeypatch.setattr(
        gerar_relatorios,
        'consultar_faturamento_mensal',
        consultar_mensal_simulado,
    )
    monkeypatch.setattr(
        gerar_relatorios,
        'consultar_produtos',
        consultar_produtos_simulado,
    )
    monkeypatch.setattr(
        gerar_relatorios,
        'exportar_excel',
        exportar_excel_simulado,
    )

    main()

    assert chamadas == [
        'resumo',
        'lojas',
        'mensal',
        'produtos',
        'exportar',
    ]